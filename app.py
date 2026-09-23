#!/usr/bin/env python3
import os,re,json,sqlite3,secrets,hashlib,hmac,csv,io,html,shutil
from openpyxl import load_workbook
from http.server import ThreadingHTTPServer,BaseHTTPRequestHandler
from urllib.parse import urlparse,parse_qs
from http.cookies import SimpleCookie
from datetime import date,datetime
import calendar
BASE=os.path.dirname(os.path.abspath(__file__))
DATA_DIR=os.environ.get('DATA_DIR', os.path.dirname(os.environ['DATABASE_PATH']) if os.environ.get('DATABASE_PATH') else BASE)
os.makedirs(DATA_DIR,exist_ok=True)
DB=os.environ.get('DATABASE_PATH',os.path.join(DATA_DIR,'eastern_gaels.db'))
SESS={}
DEMO_FILE=os.path.join(DATA_DIR,'demographics_current.json')
DEMO_SOURCE=os.path.join(BASE,'demographics.json')
def load_demo():
 p=DEMO_FILE if os.path.exists(DEMO_FILE) else DEMO_SOURCE
 with open(p,encoding='utf8') as f:d=json.load(f)
 # Normalise current and legacy demographics field names.
 d.setdefault('catchment', d.get('areas', {}))
 d.setdefault('school_totals', d.get('schools', {}))
 d.setdefault('areas', d.get('catchment', {}))
 d.setdefault('schools', d.get('school_totals', {}))
 d.setdefault('age_totals', {})
 d.setdefault('age_areas', {})
 d.setdefault('age_schools', {})
 d.setdefault('games', {})
 d.setdefault('growth', {})
 d.setdefault('players_by_age', {})
 return d
DEMO=load_demo()
COACH_FILE=os.path.join(DATA_DIR,"coaching_team_current.json")
COACH_SOURCE=os.path.join(BASE,"coaching_team.json")
def load_coaching():
 p=COACH_FILE if os.path.exists(COACH_FILE) else COACH_SOURCE
 if not os.path.exists(p): return {"coaches":{},"teams":{},"imported_at":None}
 with open(p,encoding="utf8") as f:d=json.load(f)
 d.setdefault("coaches",{});d.setdefault("teams",{});return d
COACHING=load_coaching()
def coach_key(x):
 k=" ".join(str(x or "").strip().upper().split());return {"KATE O SULLIVAN":"KATE SULLIVAN"}.get(k,k)
def parse_expiry_text(x):
 t=str(x or "").strip();m=re.search(r"(?i)expires\s+([A-Za-z]+)\s+(20\d{2})",t)
 if not m:return None
 try:
  month=datetime.strptime(m.group(1),"%B").month;year=int(m.group(2));return date(year,month,calendar.monthrange(year,month)[1]).isoformat()
 except:return None
def yn(x):
 t=str(x or '').strip().upper()
 if t in ('YES','Y','TRUE','1'): return 'Yes'
 if t in ('NO','N','FALSE','0'): return 'No'
 return 'Unknown'
def parse_coaching_xlsx(path):
 wb=load_workbook(path,data_only=True,read_only=True);ws=next((wb[n] for n in wb.sheetnames if clean_name(n).startswith('GARDA VETTING')),None)
 if ws is None:raise ValueError('Missing Garda Vetting & Safeguarding worksheet')
 coaches={}
 for row in ws.iter_rows(min_row=2,max_col=5,values_only=True):
  if not row[0]:continue
  k=coach_key(row[0]);coaches[k]={'name':' '.join(str(row[0]).strip().split()),'garda_vetted':yn(row[1]),'garda_expiry':parse_expiry_text(row[2]),'garda_expiry_text':str(row[2] or '').strip(),'safeguarding':yn(row[3]),'qualification':str(row[4] or 'Unknown').strip(),'teams':[]}
 qws=next((wb[n] for n in wb.sheetnames if clean_name(n).startswith('COACHES COURSE QUALIFICATION')),None)
 if qws:
  for row in qws.iter_rows(min_row=2,max_col=2,values_only=True):
   if not row[0]:continue
   k=coach_key(row[0]);q=str(row[1] or 'Unknown').strip()
   if k in coaches and q.upper() not in ('','UNKNOWN','NONE'):coaches[k]['qualification']=q
 pws=next((wb[n] for n in wb.sheetnames if clean_name(n).startswith('COACHES POOL SUMMARY')),None);teams={}
 if pws:
  for row in pws.iter_rows(min_row=2,values_only=True):
   if not row[0]:continue
   team=' '.join(str(row[0]).strip().upper().split());names=[]
   for raw in row[2:]:
    if not raw or not str(raw).strip():continue
    k=coach_key(raw)
    if k not in coaches:coaches[k]={'name':' '.join(str(raw).strip().title().split()),'garda_vetted':'Unknown','garda_expiry':None,'garda_expiry_text':'','safeguarding':'Unknown','qualification':'Unknown','teams':[]}
    names.append(coaches[k]['name']);coaches[k]['teams'].append(team)
   teams[team]=names
 return {'coaches':coaches,'teams':teams,'imported_at':datetime.now().isoformat(timespec='seconds')}

def expiry_bucket(c,today=None):
 today=today or date.today();v=str(c.get('garda_vetted',c.get('garda_status','Unknown'))).strip().upper()
 if v=='NO':return ('not_vetted',None)
 if v!='YES':return ('unknown',None)
 ex=c.get('garda_expiry')
 if not ex:return ('unknown_expiry',None)
 try:d=date.fromisoformat(ex);days=(d-today).days
 except:return ('unknown_expiry',None)
 if days<0:return ('expired',days)
 if days<=90:return ('urgent',days)
 if days<=180:return ('warning',days)
 return ('valid',days)


def clean_name(x):
 return ' '.join(str(x or '').strip().upper().split())
def num(x):
 try:return int(x or 0)
 except:return 0
def parse_demographics_xlsx(path):
 wb=load_workbook(path,data_only=True,read_only=True)
 names=wb.sheetnames
 def find_sheet(prefix):
  for n in names:
   if clean_name(n).startswith(clean_name(prefix)):return wb[n]
  raise ValueError('Missing worksheet: '+prefix)
 agews=find_sheet('Age Group Numbers Breakdown')
 age_totals={};age_areas={};age_schools={};warnings=[]
 for col,g in zip([1,4,7,10,13,16,19,22],['U12','U11','U10','U9','U8','U7','U6','U5']):
  areas={clean_name(agews.cell(r,col).value):num(agews.cell(r,col+1).value) for r in range(2,8) if agews.cell(r,col).value}
  schools={clean_name(agews.cell(r,col).value):num(agews.cell(r,col+1).value) for r in range(10,17) if agews.cell(r,col).value}
  total=num(agews.cell(8,col+1).value) or sum(areas.values())
  age_totals[g]=total;age_areas[g]=areas;age_schools[g]=schools
  if sum(areas.values())!=total:warnings.append(f'{g} area breakdown sums to {sum(areas.values())}, total is {total}')
  if sum(schools.values())!=total:warnings.append(f'{g} school breakdown sums to {sum(schools.values())}, total is {total}')
 def two_col(prefix,skip_totals=True):
  ws=find_sheet(prefix);d={}
  for row in ws.iter_rows(min_row=2,max_col=2,values_only=True):
   k=clean_name(row[0]);v=row[1]
   if not k or (skip_totals and k.startswith('TOTAL')):continue
   if v is None or str(v).strip()=='':continue
   d[k]=num(v)
  return d
 # Player-name worksheets: import only player name and school.
 # Discover sheets by their headers, not by worksheet name, so tabs such as
 # "U11 Players" or a default Excel name such as "Sheet10" are supported.
 # Deliberately ignore DOB, parent/guardian, phone, address and email fields.
 players_by_age={}
 for sheet_name in names:
  pws=wb[sheet_name]
  header_row=None;name_col=None;school_col=None;dob_col=None
  for r in range(1,min(11,(pws.max_row or 10)+1)):
   row_headers={}
   for c in range(1,min(21,(pws.max_column or 20)+1)):
    h=clean_name(pws.cell(r,c).value)
    row_headers[h]=c
   if 'PLAYER NAME' in row_headers and 'SCHOOL' in row_headers:
    header_row=r;name_col=row_headers['PLAYER NAME'];school_col=row_headers['SCHOOL'];dob_col=row_headers.get('DATE OF BIRTH');break
  if not (header_row and name_col and school_col):
   continue
  # First prefer an explicit Uxx / UNDER xx marker in the title or top rows.
  probe=sheet_name+' '+ ' '.join(str(v or '') for row in pws.iter_rows(min_row=1,max_row=min(6,pws.max_row or 6),max_col=min(10,pws.max_column or 10),values_only=True) for v in row)
  m=re.search(r'(?i)UNDER\s*(\d{1,2})',probe) or re.search(r'(?i)\bU\s*(\d{1,2})\b',probe)
  group=('U'+str(int(m.group(1)))) if m else None
  # If the sheet has a generic name, infer the group from DOB birth years and
  # the latest populated membership year (e.g. 2026 - 2015 = U11).
  if not group and dob_col:
   birth_years=[]
   for row in pws.iter_rows(min_row=header_row+1,max_col=dob_col,values_only=True):
    v=row[dob_col-1]
    if not v: continue
    y=None
    if hasattr(v,'year'): y=v.year
    else:
     mm=re.search(r'(19|20)\d{2}',str(v))
     if mm: y=int(mm.group(0))
    if y: birth_years.append(y)
   if birth_years:
    from collections import Counter
    birth_year=Counter(birth_years).most_common(1)[0][0]
    # Growth Year on Year is the authoritative programme year where available.
    try:
     gws=find_sheet('Growth Year on Year');years=[]
     for rr in gws.iter_rows(min_row=2,max_col=2,values_only=True):
      try:
       yy=int(float(rr[0])); total=rr[1]
       if 2000 <= yy <= 2100 and total is not None and str(total).strip(): years.append(yy)
      except (TypeError,ValueError): pass
     ref_year=max(years) if years else datetime.now().year
    except Exception:
     ref_year=datetime.now().year
    inferred=ref_year-birth_year
    if 4 <= inferred <= 18: group='U'+str(inferred)
  if not group:
   warnings.append(f'Could not determine age group for player worksheet: {sheet_name}')
   continue
  rows=[]
  for row in pws.iter_rows(min_row=header_row+1,max_col=max(name_col,school_col),values_only=True):
   player=' '.join(str(row[name_col-1] or '').strip().split())
   school=' '.join(str(row[school_col-1] or '').strip().split())
   if not player: continue
   if clean_name(player).startswith('UNDER '): continue
   rows.append({'name':player,'school':school})
  if rows: players_by_age[group]=rows

 games=two_col('Games Played By Age Group')
 # Growth sheet uses Excel numeric years (e.g. 2023.0). Parse them as years
 # instead of passing them through clean_name(), which would produce '2023.0'.
 growthws=find_sheet('Growth Year on Year');growth={}
 for yr,val in growthws.iter_rows(min_row=2,max_col=2,values_only=True):
  try:
   year=int(float(yr))
  except (TypeError,ValueError):
   continue
  if year < 1900 or year > 2200 or val is None or str(val).strip()=='':
   continue
  growth[str(year)]=num(val)
 areas=two_col('Parish Catchment Areas')
 schools=two_col('Schools Catchment Areas')
 total_players=growth.get(max(growth.keys(),key=int),sum(age_totals.values())) if growth else sum(age_totals.values())
 if sum(areas.values())!=total_players:warnings.append(f'Parish catchment sums to {sum(areas.values())}, latest player total is {total_players}')
 if sum(schools.values())!=total_players:warnings.append(f'School catchment sums to {sum(schools.values())}, latest player total is {total_players}')
 return {'age_totals':age_totals,'age_areas':age_areas,'age_schools':age_schools,'games':games,'growth':growth,'areas':areas,'schools':schools,'catchment':areas,'school_totals':schools,'players_by_age':players_by_age,'warnings':warnings,'imported_at':datetime.now().isoformat(timespec='seconds')}

def dbc(): c=sqlite3.connect(DB);c.row_factory=sqlite3.Row;return c
def hp(p,s=None):
 s=s or secrets.token_bytes(16);return s.hex()+'$'+hashlib.pbkdf2_hmac('sha256',p.encode(),s,210000).hex()
def okpw(p,x):
 try:s,h=x.split('$');return hmac.compare_digest(hp(p,bytes.fromhex(s)).split('$')[1],h)
 except:return False
def init():
 c=dbc();c.executescript("""CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,name TEXT,email TEXT UNIQUE,password TEXT,role TEXT,coach_name TEXT,active INTEGER DEFAULT 1);CREATE TABLE IF NOT EXISTS sessions(id TEXT PRIMARY KEY,date TEXT,day TEXT,school TEXT,coach TEXT,title TEXT,start TEXT,end TEXT,class_group TEXT,age_group TEXT,planned_status TEXT,actual_status TEXT DEFAULT 'Scheduled',attendance INTEGER,notes TEXT,updated_at TEXT);CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,ts TEXT,user_email TEXT,action TEXT,session_id TEXT,detail TEXT);""")
 if c.execute('select count(*) from sessions').fetchone()[0]==0:
  t=open(os.path.join(BASE,'data.js'),encoding='utf8').read();m=re.search(r'const DATA=(.*);\s*$',t,re.S)
  for s in json.loads(m.group(1)):
   c.execute('insert into sessions(id,date,day,school,coach,title,start,end,class_group,age_group,planned_status) values(?,?,?,?,?,?,?,?,?,?,?)',(s['id'],s['date'],s['day'],s['school'],s['coach'],s['session'],s['start'],s['end'],s['group'],s['age'],str(s['status']).strip()))
 c.commit();c.close()
def e(x):return html.escape(str(x or ''))
def fd(x):
 try:return datetime.strptime(x,'%Y-%m-%d').strftime('%a %d %b %Y')
 except:return x
def page(title,b,u=None):
 nav=''
 if u:
  nav='<nav><a href="/">Today</a><a href="/schedule">Schedule</a><a href="/schools">Schools</a><a href="/demographics">Demographics</a><a href="/coaching-team">Coaching Team</a><a href="/reports">Reports</a>'+'<a href="/admin">Admin</a>'+'<i></i><span>'+e(u['name'])+'</span><a href="/logout">Log out</a></nav>'
 return f'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{e(title)} · Eastern Gaels</title><style>:root{{--g:#164d2c;--bg:#f3f6f3;--m:#657269}}*{{box-sizing:border-box}}body{{font-family:system-ui;margin:0;background:var(--bg);color:#18231b}}header{{background:var(--g);color:#fff;padding:18px}}header div,main{{max-width:1100px;margin:auto}}header h1{{margin:0;font-size:22px}}nav{{display:flex;gap:5px;align-items:center;background:#fff;padding:8px max(12px,calc((100% - 1100px)/2));box-shadow:0 1px 8px #0001;overflow:auto}}nav a{{color:var(--g);font-weight:700;text-decoration:none;padding:8px}}nav i{{flex:1}}nav span{{white-space:nowrap;color:var(--m)}}main{{padding:18px}}.card{{background:#fff;padding:18px;border-radius:14px;margin:14px 0;box-shadow:0 2px 12px #0000000d}}.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:14px}}.stat{{font-size:30px;font-weight:800;color:var(--g)}}.kpi{{position:relative;overflow:hidden}}.kpi small{{display:block;color:var(--m);margin-top:3px}}.barrow{{display:grid;grid-template-columns:minmax(120px,1.4fr) 3fr 52px;gap:10px;align-items:center;margin:11px 0}}.bartrack{{height:14px;background:#edf1ed;border-radius:999px;overflow:hidden}}.barfill{{height:100%;background:var(--g);border-radius:999px;min-width:2px}}.chartlink{{color:inherit;text-decoration:none}}.progress{{height:9px;background:#e8eee9;border-radius:999px;overflow:hidden;margin-top:8px}}.progress span{{display:block;height:100%;background:var(--g)}}.donut{{width:150px;height:150px;border-radius:50%;margin:12px auto;display:grid;place-items:center;background:conic-gradient(#164d2c 0 var(--done),#b85c5c var(--done) var(--cancel),#d8a82d var(--cancel) var(--closed),#dfe5e0 var(--closed) 100%)}}.donut:after{{content:'';width:92px;height:92px;background:white;border-radius:50%}}.legend{{display:flex;gap:12px;flex-wrap:wrap;font-size:13px}}.dot{{width:10px;height:10px;border-radius:50%;display:inline-block;margin-right:5px;background:#164d2c}}.muted{{color:var(--m)}}.session{{border-left:5px solid #267345}}.badge{{display:inline-block;padding:4px 9px;border-radius:999px;background:#e5f2e8;font-size:12px;font-weight:700}}.Completed{{background:#dff2e4}}.Cancelled{{background:#f7dddd}}.Rescheduled{{background:#fff0cc}}.expired,.urgent{{background:#f7dddd;color:#8b1e1e}}.warning{{background:#fff0cc;color:#7a5700}}.unknown{{background:#eceff1;color:#4d5960}}.valid{{background:#dff2e4;color:#1d6335}}button,.btn{{display:inline-block;border:0;background:var(--g);color:#fff;padding:10px 13px;border-radius:9px;font-weight:700;text-decoration:none;cursor:pointer}}.secondary{{background:#e8eee9!important;color:#18231b!important}}input,select,textarea{{width:100%;padding:10px;border:1px solid #ccd5ce;border-radius:8px;font:inherit}}label{{display:block;font-weight:700;margin:10px 0 5px}}.row{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}.actions{{display:flex;gap:7px;flex-wrap:wrap;margin-top:12px}}table{{width:100%;border-collapse:collapse;background:#fff}}th,td{{padding:10px;border-bottom:1px solid #e4e9e5;text-align:left}}th{{font-size:12px;color:var(--m)}}.tw{{overflow:auto;border-radius:12px}}.login{{max-width:430px;margin:45px auto}}@media(max-width:650px){{main{{padding:12px}}.row{{grid-template-columns:1fr}}nav span{{display:none}}}}</style></head><body><header><div><h1>Eastern Gaels</h1><small>Coaching, Demographics & Compliance · V5</small></div></header>{nav}<main>{b}</main></body></html>'''
def card(s,admin=False):return f'''<div class="card session"><span class="badge {e(s['actual_status'])}">{e(s['actual_status'])}</span><h2>{e(s['school'])}</h2><div class="muted">{fd(s['date'])} · {e(s['start'])}–{e(s['end'])}</div><p><b>{e(s['coach'])}</b> · {e(s['class_group'])} · {e(s['age_group'])}<br>{e(s['title'])}</p><div class="actions"><a class="btn" href="/session?id={e(s['id'])}">Open session</a>{'<a class="btn secondary" href="/edit?id='+e(s['id'])+'">Edit</a>' if admin else ''}</div></div>'''
class H(BaseHTTPRequestHandler):
 def log_message(self,*a):pass
 def out(self,x,n=200,h=None):
  b=x.encode();self.send_response(n);self.send_header('Content-Type','text/html; charset=utf-8');self.send_header('Content-Length',str(len(b)));[(self.send_header(k,v)) for k,v in (h or {}).items()];self.end_headers();self.wfile.write(b)
 def red(self,x):self.send_response(303);self.send_header('Location',x);self.end_headers()
 def form(self):
  n=int(self.headers.get('Content-Length','0'));q=parse_qs(self.rfile.read(n).decode());return {k:(v if k=='ids' else v[-1]) for k,v in q.items()}
 def upload_file(self, field_name=None):
  n=int(self.headers.get('Content-Length','0'));raw=self.rfile.read(n);ct=self.headers.get('Content-Type','')
  m=re.search(r'boundary=(?:\"([^\"]+)\"|([^;]+))',ct)
  if not m:return None,None
  boundary=(m.group(1) or m.group(2)).strip().encode();parts=raw.split(b'--'+boundary)
  wanted=(('name="'+field_name+'"').encode() if field_name else None)
  for part in parts:
   head,sep,body=part.partition(b'\r\n\r\n')
   if not sep or b'filename=' not in head:continue
   if wanted and wanted not in head:continue
   fm=re.search(br'filename="([^"]*)"',head)
   name=fm.group(1).decode('utf-8',errors='ignore') if fm else 'upload.xlsx'
   # A multipart part ends with one CRLF before the next boundary. Remove only
   # that framing CRLF; rstrip() can corrupt the binary tail of an .xlsx ZIP.
   if body.endswith(b'\r\n'):body=body[:-2]
   return os.path.basename(name.replace('\\','/')),body
  return None,None
 def user(self):
  z=SimpleCookie(self.headers.get('Cookie'));sid=z.get('egsid');uid=SESS.get(sid.value) if sid else None
  if not uid:return None
  c=dbc();u=c.execute('select * from users where id=? and active=1',(uid,)).fetchone();c.close();return u
 def need(self):
  u=self.user()
  if not u:self.red('/login')
  return u
 def vis(self,u):return (' ',[])
 def do_GET(self):
  p=urlparse(self.path);path=p.path;q=parse_qs(p.query);c=dbc()
  if path=='/setup':
   if c.execute('select count(*) from users').fetchone()[0]:c.close();return self.red('/login')
   c.close();return self.out(page('Setup','<div class="card login"><h2>Create administrator</h2><p class="muted">First-run setup.</p><form method="post"><label>Name</label><input name="name" required><label>Email</label><input name="email" type="email" required><label>Password</label><input name="password" type="password" minlength="8" required><div class="actions"><button>Create admin</button></div></form></div>'))
  if path=='/login':
   if not c.execute('select count(*) from users').fetchone()[0]:c.close();return self.red('/setup')
   c.close();return self.out(page('Login','<div class="card login"><h2>Log in</h2><form method="post"><label>Email</label><input name="email" type="email" required><label>Password</label><input name="password" type="password" required><div class="actions"><button>Log in</button></div></form></div>'))
  if path=='/logout':
   z=SimpleCookie(self.headers.get('Cookie'));sid=z.get('egsid');SESS.pop(sid.value,None) if sid else None;c.close();self.send_response(303);self.send_header('Location','/login');self.send_header('Set-Cookie','egsid=; Path=/; Max-Age=0; HttpOnly; SameSite=Lax');self.end_headers();return
  u=self.need()
  if not u:c.close();return
  if path=='/':
   d=q.get('date',[date.today().isoformat()])[0];w,a=self.vis(u);con='and' if 'where' in w else 'where';r=c.execute('select * from sessions '+w+con+' date=? order by start',a+[d]).fetchall();c.close();b=f'<h2>Today / Selected Day</h2><form class="card"><label>Date</label><div class="row"><input type="date" name="date" value="{e(d)}"><button>Show date</button></div></form>'+(''.join(card(x,True) for x in r) if r else '<div class="card">No coaching scheduled for this date.</div>');return self.out(page('Today',b,u))
  if path=='/schedule':
   w,a=self.vis(u);r=c.execute('select * from sessions '+w+' order by date,start',a).fetchall();c.close();trs=''.join(f'<tr><td>{fd(x["date"])}</td><td><a href="/session?id={e(x["id"])}">{e(x["school"])}</a></td><td>{e(x["coach"])}</td><td>{e(x["class_group"])}</td><td><span class="badge {e(x["actual_status"])}">{e(x["actual_status"])}</span></td></tr>' for x in r);b=f'''<h2>Full Schedule</h2><div class="card"><input id="s" placeholder="Search school, coach, class or date…" oninput="f()"></div><div class="tw"><table id="t"><tr><th>Date</th><th>School</th><th>Coach</th><th>Class</th><th>Status</th></tr>{trs}</table></div><script>function f(){{let q=s.value.toLowerCase();document.querySelectorAll('#t tr').forEach((r,i)=>{{if(i)r.style.display=r.innerText.toLowerCase().includes(q)?'':'none'}})}}</script>''';return self.out(page('Schedule',b,u))
  if path=='/schools':
   w,a=self.vis(u);r=c.execute('select school,count(*) n,sum(actual_status="Completed") done from sessions '+w+' group by school order by school',a).fetchall();c.close();return self.out(page('Schools','<h2>Schools</h2><div class="grid">'+''.join(f'<div class="card"><h3>{e(x["school"])}</h3><div class="stat">{x["done"]}</div><div class="muted">completed of {x["n"]} scheduled</div></div>' for x in r)+'</div>',u))
  if path=='/session':
   s=c.execute('select * from sessions where id=?',(q.get('id',[''])[0],)).fetchone();c.close()
   if not s:return self.out(page('Not found','<div class="card">Session not found.</div>',u),404)
   opts=''.join('<option'+(' selected' if s['actual_status']==x else '')+'>'+x+'</option>' for x in ['Scheduled','Completed','Cancelled','School Closed','Rescheduled']);b=card(s,True)+f'''<div class="card"><form method="post"><input type="hidden" name="id" value="{e(s['id'])}"><label>Status</label><select name="status">{opts}</select><label>Attendance / children coached</label><input type="number" min="0" name="attendance" value="{e(s['attendance'])}"><label>Session notes</label><textarea name="notes" rows="5">{e(s['notes'])}</textarea><div class="actions"><button>Save session record</button></div></form></div>''';return self.out(page('Session',b,u))
  if path=='/demographics':
   c.close();years=sorted((int(y) for y in DEMO.get('growth',{}) if str(y).isdigit()));latest_year=years[-1] if years else datetime.now().year;prev_year=years[-2] if len(years)>1 else latest_year-1;total=DEMO.get('growth',{}).get(str(latest_year),sum(DEMO.get('age_totals',{}).values()));prev=DEMO.get('growth',{}).get(str(prev_year),0);growthpct=round((total-prev)*100/prev) if prev else 0;games=sum(DEMO.get('games',{}).values())
   def bars(data):
    mx=max([int(v) for v in data.values()] or [1]);return ''.join(f'<div class="barrow"><b>{e(k)}</b><div class="bartrack"><div class="barfill" style="width:{int(v)*100/mx:.0f}%"></div></div><b>{int(v)}</b></div>' for k,v in data.items())
   age=dict(sorted(DEMO['age_totals'].items(),key=lambda x:int(x[0][1:])))
   agelinks=''.join(f'<a class="chartlink" href="/demographics/age?group={e(k)}"><div class="card"><h3>{e(k)}</h3><div class="stat">{int(v)}</div><div class="muted">players · tap for catchment & schools</div></div></a>' for k,v in age.items())
   b=f'''<h2>Club Demographics</h2><p class="muted">Eastern Gaels player profile · {latest_year}</p><div class="grid"><div class="card kpi"><div class="stat">{total}</div><b>Players</b><small>{latest_year} club population</small></div><div class="card kpi"><div class="stat">+{growthpct}%</div><b>Year-on-year growth</b><small>{prev} players in {prev_year}</small></div><div class="card kpi"><div class="stat">{len(DEMO['age_totals'])}</div><b>Age groups</b><small>U5 through U12</small></div><div class="card kpi"><div class="stat">{games}</div><b>Games played</b><small>recorded in source workbook</small></div></div><div class="grid"><div class="card"><h3>Membership growth</h3>{bars(DEMO['growth'])}</div><div class="card"><h3>Players by age group</h3>{bars(age)}</div></div><div class="grid"><div class="card"><h3>Where our players live</h3>{bars(DEMO['catchment'])}</div><div class="card"><h3>Schools represented</h3>{bars(DEMO['school_totals'])}</div></div><div class="card"><h3>Games played by team</h3>{bars(DEMO['games'])}</div><h3>Age-group detail</h3><div class="grid">{agelinks}</div>''';return self.out(page('Club Demographics',b,u))
  if path=='/demographics/age':
   c.close();g=q.get('group',[''])[0].upper()
   if g not in DEMO['age_totals']:return self.red('/demographics')
   def bars2(data):
    mx=max([int(v) for v in data.values()] or [1]);return ''.join(f'<div class="barrow"><b>{e(k)}</b><div class="bartrack"><div class="barfill" style="width:{int(v)*100/mx:.0f}%"></div></div><b>{int(v)}</b></div>' for k,v in data.items())
   players=DEMO.get('players_by_age',{}).get(g,[]);prows=''.join(f'<tr><td>{e(x.get("name",""))}</td><td>{e(x.get("school",""))}</td></tr>' for x in players);plist=(f'<div class="card"><h3>{e(g)} Players</h3><p class="muted">{len(players)} player names imported. Only player name and school are stored/displayed.</p><div class="tw"><table><tr><th>Player Name</th><th>School</th></tr>{prows}</table></div></div>' if players else '<div class="card"><h3>Players</h3><p class="muted">No player-name list has been imported for this age group yet.</p></div>');total=int(DEMO['age_totals'][g]);b=f'''<div class="actions"><a class="btn secondary" href="/demographics">← Demographics</a></div><h2>{e(g)} Demographics</h2><div class="card kpi"><div class="stat">{total}</div><b>{e(g)} players</b><small>2026</small></div>{plist}<div class="grid"><div class="card"><h3>Residential catchment</h3>{bars2(DEMO['age_areas'][g])}</div><div class="card"><h3>School breakdown</h3>{bars2(DEMO['age_schools'][g])}</div></div><div class="card"><p class="muted">Figures are imported from the updated Eastern Gaels demographics workbook. Zero-value categories are retained so the source structure remains visible.</p></div>''';return self.out(page(g+' Demographics',b,u))
  if path=='/coaching-team':
   c.close();coaches=list(COACHING.get('coaches',{}).values());teams=COACHING.get('teams',{});buckets={'expired':0,'urgent':0,'warning':0,'unknown':0,'unknown_expiry':0,'not_vetted':0,'valid':0}
   for x in coaches:buckets[expiry_bucket(x)[0]]+=1
   vetted=sum(str(x.get('garda_vetted',x.get('garda_status',''))).strip().upper()=='YES' for x in coaches);safe=sum(str(x.get('safeguarding','')).strip().upper()=='YES' for x in coaches);qual=sum(str(x.get('qualification','')).strip().upper() not in ('','UNKNOWN','NONE','NO') for x in coaches);rows='';order={'expired':0,'urgent':1,'warning':2,'unknown_expiry':3,'not_vetted':4,'unknown':5,'valid':6}
   labels={'expired':'Expired','urgent':'Expires <=90 days','warning':'Expires <=180 days','unknown_expiry':'Vetted - expiry missing','not_vetted':'Not vetted','unknown':'Vetting unknown','valid':'Vetted / valid'}
   for x in sorted(coaches,key=lambda z:(order[expiry_bucket(z)[0]],z.get('name',''))):
    bucket,days=expiry_bucket(x);expiry=x.get('garda_expiry_text') or x.get('garda_expiry') or '-';detail=('Expired '+str(abs(days))+' days ago' if bucket=='expired' and days is not None else (str(days)+' days remaining' if days is not None else expiry));sv=str(x.get('safeguarding','Unknown'));qv=str(x.get('qualification','Unknown'))
    rows+=f'<tr><td><b>{e(x.get("name"))}</b></td><td>{e(x.get("garda_vetted",x.get("garda_status","Unknown")))}</td><td><span class="badge {bucket}">{labels[bucket]}</span><br><small>{e(detail)}</small></td><td>{e(sv)}</td><td>{e(qv)}</td><td>{e(", ".join(x.get("teams",[])))}</td></tr>'
   teamcards=''.join(f'<div class="card"><h3>{e(k)}</h3><div class="stat">{len(v)}</div><div class="muted">coaches</div><p>{e(", ".join(v))}</p></div>' for k,v in teams.items())
   b=f'''<h2>Coaching Team & Compliance</h2><p class="muted">Garda Vetting and Safeguarding are tracked independently. Vetting expiry warnings apply only when Garda Vetted is Yes.</p><div class="grid"><div class="card kpi"><div class="stat">{len(coaches)}</div><b>Total coaches</b></div><div class="card kpi"><div class="stat">{vetted}</div><b>Garda vetted</b><small>of {len(coaches)}</small></div><div class="card kpi"><div class="stat">{safe}</div><b>Safeguarding completed</b><small>of {len(coaches)}</small></div><div class="card kpi"><div class="stat">{qual}</div><b>Qualified coaches</b></div><div class="card kpi"><div class="stat">{buckets['expired']+buckets['urgent']}</div><b>Vetting expiry action</b><small>expired / within 90 days</small></div></div><div class="card"><h3>Coach compliance</h3><div class="tw"><table><tr><th>Coach</th><th>Garda Vetted</th><th>Vetting Expiry</th><th>Safeguarding</th><th>Qualification</th><th>Teams</th></tr>{rows}</table></div></div><h3>Coaching pool</h3><div class="grid">{teamcards}</div>'''
   return self.out(page('Coaching Team',b,u))
  if path=='/reports':
   w,a=self.vis(u)
   ss=c.execute('select * from sessions '+w+' order by date,start',a).fetchall()
   total=len(ss);done=sum(x['actual_status']=='Completed' for x in ss);cancelled=sum(x['actual_status']=='Cancelled' for x in ss);closed=sum(x['actual_status']=='School Closed' for x in ss);scheduled=sum(x['actual_status']=='Scheduled' for x in ss);attendance=sum((x['attendance'] or 0) for x in ss if x['actual_status']=='Completed')
   hours=0.0
   for x in ss:
    if x['actual_status']=='Completed' and x['start'] and x['end']:
     try:
      st=datetime.strptime(x['start'],'%H:%M');en=datetime.strptime(x['end'],'%H:%M');hours+=max(0,(en-st).seconds/3600)
     except:pass
   pct=round(done*100/total) if total else 0
   byschool={};bycoach={};bymonth={}
   for x in ss:
    for dct,key in ((byschool,x['school']),(bycoach,x['coach'])):
     z=dct.setdefault(key,[0,0]);z[1]+=1;z[0]+=x['actual_status']=='Completed'
    if x['actual_status']=='Completed':
     try:key=datetime.strptime(x['date'],'%Y-%m-%d').strftime('%b %Y')
     except:key=x['date'][:7]
     bymonth[key]=bymonth.get(key,0)+1
   maxschool=max(1,max([v[0] for v in byschool.values()] or [0]));maxcoach=max(1,max([v[0] for v in bycoach.values()] or [0]));maxmonth=max(1,max(list(bymonth.values()) or [0]))
   schoolbars=''.join(f'<a class="chartlink" href="/reports/detail?school={e(k)}"><div class="barrow"><b>{e(k)}</b><div class="bartrack"><div class="barfill" style="width:{v[0]*100/maxschool:.0f}%"></div></div><b>{v[0]}</b></div></a>' for k,v in sorted(byschool.items()))
   coachbars=''.join(f'<a class="chartlink" href="/reports/detail?coach={e(k)}"><div class="barrow"><b>{e(k)}</b><div class="bartrack"><div class="barfill" style="width:{v[0]*100/maxcoach:.0f}%"></div></div><b>{v[0]}</b></div></a>' for k,v in sorted(bycoach.items()))
   monthbars=''.join(f'<div class="barrow"><b>{e(k)}</b><div class="bartrack"><div class="barfill" style="width:{v*100/maxmonth:.0f}%"></div></div><b>{v}</b></div>' for k,v in bymonth.items()) or '<p class="muted">No completed sessions yet.</p>'
   schoolcards=''.join(f'<a class="chartlink" href="/reports/detail?school={e(k)}"><div class="card"><h3>{e(k)}</h3><div class="stat">{v[0]} / {v[1]}</div><div class="muted">sessions completed · {round(v[0]*100/v[1]) if v[1] else 0}%</div><div class="progress"><span style="width:{v[0]*100/v[1] if v[1] else 0:.0f}%"></span></div></div></a>' for k,v in sorted(byschool.items()))
   d1=pct;d2=min(100,d1+(cancelled*100/total if total else 0));d3=min(100,d2+(closed*100/total if total else 0))
   c.close();b=f'''<h2>Coaching Dashboard</h2><p class="muted">Live summary of the 2026–2027 schools coaching programme.</p><div class="grid"><div class="card kpi"><div class="stat">{done}</div><b>Sessions completed</b><small>of {total} planned</small></div><div class="card kpi"><div class="stat">{pct}%</div><b>Programme complete</b><small>{scheduled} still scheduled</small></div><div class="card kpi"><div class="stat">{hours:.1f}</div><b>Coaching hours</b><small>completed sessions</small></div><div class="card kpi"><div class="stat">{attendance}</div><b>Recorded attendance</b><small>total contacts entered</small></div></div><div class="grid"><div class="card"><h3>Completed sessions by month</h3>{monthbars}</div><div class="card"><h3>Session status</h3><div class="donut" style="--done:{d1}%;--cancel:{d2}%;--closed:{d3}%"></div><div class="legend"><span><i class="dot"></i>Completed {done}</span><span>Cancelled {cancelled}</span><span>School closed {closed}</span><span>Scheduled {scheduled}</span></div></div></div><div class="grid"><div class="card"><h3>Completed by school</h3>{schoolbars}</div><div class="card"><h3>Completed by coach</h3>{coachbars}</div></div><h3>School progress</h3><div class="grid">{schoolcards}</div><div class="actions"><a class="btn" href="/export.csv">Download CSV report</a></div>''';return self.out(page('Dashboard',b,u))
  if path=='/reports/detail':
   school=q.get('school',[''])[0];coach=q.get('coach',[''])[0];w,a=self.vis(u);con='and' if 'where' in w else 'where';extra='';vals=list(a);title='Completed sessions'
   if school:extra=' school=?';vals.append(school);title=school
   elif coach:extra=' coach=?';vals.append(coach);title=coach
   else:extra=" actual_status='Completed'"
   r=c.execute('select * from sessions '+w+con+extra+' order by date,start',vals).fetchall();c.close();rows=''.join(f'<tr><td>{fd(x["date"])}</td><td>{e(x["school"])}</td><td>{e(x["coach"])}</td><td>{e(x["class_group"])}</td><td><span class="badge {e(x["actual_status"])}">{e(x["actual_status"])}</span></td><td>{e(x["attendance"])}</td></tr>' for x in r);return self.out(page('Report detail',f'<h2>{e(title)}</h2><p class="muted">{len(r)} sessions</p><div class="tw"><table><tr><th>Date</th><th>School</th><th>Coach</th><th>Class</th><th>Status</th><th>Attendance</th></tr>{rows}</table></div><div class="actions"><a class="btn secondary" href="/reports">Back to dashboard</a></div>',u))
  if path=='/admin':
   us=c.execute('select * from users order by role,name').fetchall();co=[x[0] for x in c.execute('select distinct coach from sessions order by coach')];now=datetime.now();today=now.date().isoformat();clock=now.strftime('%H:%M');overdue=c.execute("select count(*) from sessions where actual_status='Scheduled' and (date < ? or (date = ? and coalesce(end,start,'23:59') < ?))",(today,today,clock)).fetchone()[0];c.close();opts=''.join(f'<option>{e(x)}</option>' for x in co);rows=''.join(f'<tr><td>{e(x["name"])}</td><td>{e(x["email"])}</td><td>{'Super Admin' if x['role']=='admin' else 'Full Access'}</td><td>{e(x["coach_name"])}</td></tr>' for x in us);b=f'''<h2>Administration</h2><div class="card"><h3>Past due sessions</h3><div class="stat">{overdue}</div><p class="muted">Scheduled sessions whose date/time has passed. Cancelled, School Closed, Rescheduled and already Completed sessions are not included.</p><div class="actions"><a class="btn" href="/admin/past-due">Review past sessions</a></div></div><div class="grid">{('<div class="card"><h3>User Management</h3><p class="muted">Super Admin only. New users receive Full Access to operational features.</p><form method="post" action="/admin/user"><label>Name</label><input name="name" required><label>Email</label><input type="email" name="email" required><label>Temporary password</label><input type="password" minlength="8" name="password" required><label>Timetable coach (optional)</label><select name="coach_name"><option value="">— None —</option>'+opts+'</select><div class="actions"><button>Create Full Access user</button></div></form></div>') if u['role']=='admin' else ''}<div class="card"><h3>Add session</h3><form method="post" action="/admin/session"><label>Date</label><input type="date" name="date" required><label>School</label><input name="school" required><label>Coach</label><input name="coach" required><div class="row"><div><label>Start</label><input type="time" name="start"></div><div><label>End</label><input type="time" name="end"></div></div><label>Class group</label><input name="class_group"><label>Age group</label><input name="age_group"><label>Title</label><input name="title" value="GAA Coaching"><div class="actions"><button>Add session</button></div></form></div></div><div class="card"><h3>Update Demographics</h3><p class="muted">Upload the latest Eastern Gaels demographics Excel workbook. The app validates it and shows a preview before anything is changed.</p><form method="post" action="/admin/demographics/preview" enctype="multipart/form-data"><label>Demographics workbook (.xlsx)</label><input type="file" name="demographics_file" accept=".xlsx" required><div class="actions"><button>Upload & preview</button></div></form></div><div class="card"><h3>Update Coaching Team</h3><p class="muted">Upload the latest Garda vetting, safeguarding, qualifications and coaching pool workbook.</p><form method="post" action="/admin/coaching-team/upload" enctype="multipart/form-data"><label>Coaching details workbook (.xlsx)</label><input type="file" name="coaching_file" accept=".xlsx" required><div class="actions"><button>Upload coaching details</button></div></form></div>{('<div class="card"><h3>Users</h3><div class="tw"><table><tr><th>Name</th><th>Email</th><th>Access</th><th>Timetable coach</th></tr>'+rows+'</table></div></div>') if u['role']=='admin' else ''}''';return self.out(page('Admin',b,u))
  if path=='/admin/past-due':
   now=datetime.now();today=now.date().isoformat();clock=now.strftime('%H:%M');r=c.execute("select * from sessions where actual_status='Scheduled' and (date < ? or (date = ? and coalesce(end,start,'23:59') < ?)) order by date,start",(today,today,clock)).fetchall();c.close()
   rows=''.join(f'<tr><td><input form="bulk" type="checkbox" name="ids" value="{e(x["id"])}" checked></td><td>{fd(x["date"])}</td><td>{e(x["school"])}</td><td>{e(x["coach"])}</td><td>{e(x["start"])}–{e(x["end"])}</td><td><a href="/session?id={e(x["id"])}">Open</a></td></tr>' for x in r)
   if not r:b='<h2>Past due sessions</h2><div class="card"><h3>All caught up</h3><p>There are no past-due sessions still marked Scheduled.</p><a class="btn secondary" href="/admin">Back to Admin</a></div>'
   else:b=f'''<h2>Past due sessions</h2><div class="card"><p>Review the sessions below. Untick anything that was not actually completed, then mark the selected sessions Completed.</p><form id="bulk" method="post" action="/admin/past-due"><div class="actions"><button>Mark selected as Completed</button><button type="button" class="secondary" onclick="document.querySelectorAll('input[name=ids]').forEach(x=>x.checked=true)">Select all</button><button type="button" class="secondary" onclick="document.querySelectorAll('input[name=ids]').forEach(x=>x.checked=false)">Clear</button></div></form></div><div class="tw"><table><tr><th>Complete?</th><th>Date</th><th>School</th><th>Coach</th><th>Time</th><th></th></tr>{rows}</table></div>'''
   return self.out(page('Past due sessions',b,u))
  if path=='/admin/demographics/preview':
   c.close();pending=os.path.join(DATA_DIR,'demographics_pending.json')
   if not os.path.exists(pending):return self.red('/admin')
   nd=json.load(open(pending,encoding='utf8'));old=load_demo();latest=lambda d:max((int(y) for y in d.get('growth',{}) if str(y).isdigit()),default=2026);ny=latest(nd);oy=latest(old);nt=nd.get('growth',{}).get(str(ny),sum(nd.get('age_totals',{}).values()));ot=old.get('growth',{}).get(str(oy),sum(old.get('age_totals',{}).values()));
   changes=''.join(f'<tr><td>{e(g)}</td><td>{old.get("age_totals",{}).get(g,0)}</td><td>{nd.get("age_totals",{}).get(g,0)}</td><td>{nd.get("age_totals",{}).get(g,0)-old.get("age_totals",{}).get(g,0):+d}</td></tr>' for g in ['U5','U6','U7','U8','U9','U10','U11','U12'])
   warns=''.join(f'<li>{e(x)}</li>' for x in nd.get('warnings',[])) or '<li>No validation warnings.</li>'
   b=f'''<h2>Preview demographics import</h2><div class="grid"><div class="card kpi"><div class="stat">{ot}</div><b>Current players</b><small>{oy}</small></div><div class="card kpi"><div class="stat">{nt}</div><b>Uploaded players</b><small>{ny}</small></div><div class="card kpi"><div class="stat">{nt-ot:+d}</div><b>Change</b><small>players</small></div></div><div class="card"><h3>Age-group changes</h3><div class="tw"><table><tr><th>Age</th><th>Current</th><th>Uploaded</th><th>Change</th></tr>{changes}</table></div></div><div class="card"><h3>Validation</h3><ul>{warns}</ul><p class="muted">Warnings do not block import; review them before confirming.</p></div><div class="card"><form method="post" action="/admin/demographics/confirm"><div class="actions"><button>Confirm import</button><a class="btn secondary" href="/admin">Cancel</a></div></form></div>'''
   return self.out(page('Preview demographics',b,u))
  if path=='/edit':
   s=c.execute('select * from sessions where id=?',(q.get('id',[''])[0],)).fetchone();c.close()
   if not s:return self.red('/schedule')
   b=f'''<h2>Edit session</h2><div class="card"><form method="post"><input type="hidden" name="id" value="{e(s['id'])}"><label>Date</label><input type="date" name="date" value="{e(s['date'])}" required><label>School</label><input name="school" value="{e(s['school'])}" required><label>Coach</label><input name="coach" value="{e(s['coach'])}" required><div class="row"><div><label>Start</label><input type="time" name="start" value="{e(s['start'])}"></div><div><label>End</label><input type="time" name="end" value="{e(s['end'])}"></div></div><label>Class</label><input name="class_group" value="{e(s['class_group'])}"><label>Age group</label><input name="age_group" value="{e(s['age_group'])}"><label>Title</label><input name="title" value="{e(s['title'])}"><div class="actions"><button>Save changes</button></div></form></div>''';return self.out(page('Edit',b,u))
  if path=='/export.csv':
   r=c.execute('select * from sessions order by date,start').fetchall();c.close();o=io.StringIO();w=csv.writer(o);w.writerow(r[0].keys() if r else []);[w.writerow(tuple(x)) for x in r];b=o.getvalue().encode();self.send_response(200);self.send_header('Content-Type','text/csv');self.send_header('Content-Disposition','attachment; filename="eastern-gaels-report.csv"');self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b);return
  c.close();return self.out(page('Not found','<div class="card">Page not found.</div>',u),404)
 def do_POST(self):
  p=urlparse(self.path).path
  if p=='/admin/coaching-team/upload':
   u=self.need()
   if not u:return self.out('Forbidden',403)
   name,data=self.upload_file('coaching_file')
   if not data or not name.lower().endswith('.xlsx'):return self.out(page('Upload error','<div class="card"><h2>Please choose a valid .xlsx file.</h2><a class="btn" href="/admin">Back</a></div>',u),400)
   tmp=os.path.join(DATA_DIR,'coaching_team_upload.xlsx')
   try:
    open(tmp,'wb').write(data);parsed=parse_coaching_xlsx(tmp);json.dump(parsed,open(COACH_FILE,'w',encoding='utf8'),indent=2);globals()['COACHING']=load_coaching()
   except Exception as ex:return self.out(page('Upload error',f'<div class="card"><h2>Could not read coaching workbook</h2><p>{e(ex)}</p><a class="btn" href="/admin">Back</a></div>',u),400)
   return self.red('/coaching-team')
  if p=='/admin/demographics/preview':
   u=self.need()
   if not u:return self.out('Forbidden',403)
   name,data=self.upload_file('demographics_file')
   if not data or not name.lower().endswith('.xlsx'):return self.out(page('Upload error','<div class="card"><h2>Please choose a valid .xlsx file.</h2><a class="btn" href="/admin">Back</a></div>',u),400)
   tmp=os.path.join(DATA_DIR,'demographics_upload.xlsx')
   try:
    open(tmp,'wb').write(data);parsed=parse_demographics_xlsx(tmp);json.dump(parsed,open(os.path.join(DATA_DIR,'demographics_pending.json'),'w',encoding='utf8'),indent=2)
   except Exception as ex:
    return self.out(page('Upload error',f'<div class="card"><h2>Could not read workbook</h2><p>{e(ex)}</p><a class="btn" href="/admin">Back</a></div>',u),400)
   return self.red('/admin/demographics/preview')
  f=self.form();c=dbc()
  if p=='/setup':
   if c.execute('select count(*) from users').fetchone()[0]:c.close();return self.red('/login')
   c.execute('insert into users(name,email,password,role) values(?,?,?,?)',(f['name'],f['email'].lower(),hp(f['password']),'admin'));c.commit();c.close();return self.red('/login')
  if p=='/login':
   u=c.execute('select * from users where email=? and active=1',(f.get('email','').lower(),)).fetchone();c.close()
   if not u or not okpw(f.get('password',''),u['password']):return self.out(page('Login','<div class="card login"><h2>Incorrect email or password</h2><a class="btn" href="/login">Try again</a></div>'),401)
   sid=secrets.token_urlsafe(32);SESS[sid]=u['id'];self.send_response(303);self.send_header('Location','/');secure='; Secure' if os.environ.get('COOKIE_SECURE','0')=='1' else ''
   self.send_header('Set-Cookie',f'egsid={sid}; Path=/; HttpOnly; SameSite=Lax{secure}');self.end_headers();return
  u=self.need()
  if not u:c.close();return
  if p=='/admin/demographics/confirm':
   pending=os.path.join(DATA_DIR,'demographics_pending.json')
   if os.path.exists(pending):
    hist=os.path.join(DATA_DIR,'demographics_history');os.makedirs(hist,exist_ok=True)
    if os.path.exists(DEMO_FILE):shutil.copy2(DEMO_FILE,os.path.join(hist,'demographics_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.json'))
    else:shutil.copy2(DEMO_SOURCE,os.path.join(hist,'demographics_initial_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.json'))
    os.replace(pending,DEMO_FILE);global DEMO;DEMO=load_demo()
   c.close();return self.red('/demographics')
  if p=='/session':
   s=c.execute('select * from sessions where id=?',(f.get('id'),)).fetchone()
   if not s:c.close();return self.out('Forbidden',403)
   a=f.get('attendance','');a=int(a) if a.isdigit() else None;c.execute('update sessions set actual_status=?,attendance=?,notes=?,updated_at=? where id=?',(f.get('status'),a,f.get('notes',''),datetime.now().isoformat(timespec='seconds'),f['id']));c.execute('insert into audit(ts,user_email,action,session_id,detail) values(?,?,?,?,?)',(datetime.now().isoformat(timespec='seconds'),u['email'],'update_session',f['id'],f.get('status')));c.commit();c.close();return self.red('/session?id='+f['id'])
  if p=='/admin/past-due':
   ids=f.get('ids',[]);ids=[ids] if isinstance(ids,str) else ids;now=datetime.now();today=now.date().isoformat();clock=now.strftime('%H:%M');changed=0
   for sid in ids:
    r=c.execute("select * from sessions where id=? and actual_status='Scheduled' and (date < ? or (date = ? and coalesce(end,start,'23:59') < ?))",(sid,today,today,clock)).fetchone()
    if r:
     c.execute("update sessions set actual_status='Completed',updated_at=? where id=?",(now.isoformat(timespec='seconds'),sid));c.execute('insert into audit(ts,user_email,action,session_id,detail) values(?,?,?,?,?)',(now.isoformat(timespec='seconds'),u['email'],'bulk_complete_past_due',sid,'Completed'));changed+=1
   c.commit();c.close();return self.red('/admin/past-due')
  if p=='/admin/user' and u['role']=='admin':
   try:c.execute('insert into users(name,email,password,role,coach_name) values(?,?,?,?,?)',(f['name'],f['email'].lower(),hp(f['password']),'full',f.get('coach_name') or None));c.commit()
   except sqlite3.IntegrityError:pass
   c.close();return self.red('/admin')
  if p=='/admin/session':
   sid='s'+secrets.token_hex(5);day=datetime.strptime(f['date'],'%Y-%m-%d').strftime('%A');c.execute('insert into sessions(id,date,day,school,coach,title,start,end,class_group,age_group,planned_status) values(?,?,?,?,?,?,?,?,?,?,?)',(sid,f['date'],day,f['school'],f['coach'],f.get('title'),f.get('start'),f.get('end'),f.get('class_group'),f.get('age_group'),'Confirmed'));c.commit();c.close();return self.red('/session?id='+sid)
  if p=='/edit':
   day=datetime.strptime(f['date'],'%Y-%m-%d').strftime('%A');c.execute('update sessions set date=?,day=?,school=?,coach=?,start=?,end=?,class_group=?,age_group=?,title=?,updated_at=? where id=?',(f['date'],day,f['school'],f['coach'],f.get('start'),f.get('end'),f.get('class_group'),f.get('age_group'),f.get('title'),datetime.now().isoformat(timespec='seconds'),f['id']));c.commit();c.close();return self.red('/session?id='+f['id'])
  c.close();return self.red('/')
if __name__=='__main__':init();port=int(os.environ.get('PORT','8000'));print(f'Open http://localhost:{port}');ThreadingHTTPServer(('0.0.0.0',port),H).serve_forever()
